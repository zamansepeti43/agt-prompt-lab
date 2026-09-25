#!/usr/bin/env python3
import csv, io, json, re, urllib.request, unicodedata, subprocess, sys
from collections import Counter, defaultdict
from pathlib import Path

CSV_URL = "https://raw.githubusercontent.com/f/prompts.chat/main/prompts.csv"
API_URL = "https://prompts.chat/api/prompts?perPage=1&page=1"

def get(url):
    req = urllib.request.Request(url, headers={"User-Agent": "agt-prompt-lab-analysis/2.0"})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read()

def clean_excel(s):
    s = str(s or "")
    return "".join(ch for ch in s if ord(ch) >= 32 or ch in "\\t\\n\\r")

def norm(s):
    s = unicodedata.normalize("NFKC", s or "").lower()
    return re.sub(r"\s+", " ", s).strip()

def provider(model):
    m = model.lower()
    if any(x in m for x in ("gpt", "o1", "o3", "o4", "openai")): return "OpenAI"
    if any(x in m for x in ("claude", "anthropic")): return "Anthropic"
    if any(x in m for x in ("gemini", "gemma", "google")): return "Google"
    if any(x in m for x in ("llama", "meta")): return "Meta"
    if any(x in m for x in ("mistral", "mixtral")): return "Mistral"
    if any(x in m for x in ("qwen", "alibaba")): return "Alibaba/Qwen"
    if "deepseek" in m: return "DeepSeek"
    if any(x in m for x in ("grok", "xai")): return "xAI"
    return "Other/Unknown"

def find_field(fields, names):
    low = {f.lower(): f for f in fields}
    for n in names:
        if n.lower() in low: return low[n.lower()]
    return None

print("Downloading CSV...")
raw = get(CSV_URL)
print(f"Downloaded {len(raw):,} bytes")
csv.field_size_limit(20 * 1024 * 1024)
text = raw.decode("utf-8-sig", errors="replace")
reader = csv.DictReader(io.StringIO(text))
rows = list(reader)
fields = reader.fieldnames or []
if not fields or not rows:
    raise RuntimeError("CSV parsed but contains no fields/rows")

title_f = find_field(fields, ["title","name"])
content_f = find_field(fields, ["content","prompt","text"])
category_f = find_field(fields, ["category","category_name","categoryName"])
type_f = find_field(fields, ["type","prompt_type","promptType"])
model_f = find_field(fields, ["bestWithModels","best_with_models","models","model","supported_models","supportedModels"])
tags_f = find_field(fields, ["tags","tag"])
id_f = find_field(fields, ["id","prompt_id","promptId"])

report = {
    "source": CSV_URL, "csv_bytes": len(raw), "csv_rows": len(rows),
    "columns": fields, "live_api": {}, "counts": {}, "model_metadata": {},
    "quality": {}, "samples": []
}

try:
    api = json.loads(get(API_URL).decode("utf-8"))
    report["live_api"] = {
        "url": API_URL, "total": api.get("total"), "page": api.get("page"),
        "perPage": api.get("perPage"), "totalPages": api.get("totalPages"), "status": "ok"
    }
except Exception as e:
    report["live_api"] = {"url": API_URL, "status": "error", "error": str(e)}

for label, f in [("category", category_f), ("type", type_f)]:
    if f:
        vals = Counter((r.get(f) or "").strip() for r in rows)
        vals.pop("", None)
        report["counts"][label] = vals.most_common(200)

models, providers = Counter(), Counter()
model_rows = 0
if model_f:
    for r in rows:
        rawv = (r.get(model_f) or "").strip()
        if not rawv: continue
        model_rows += 1
        try:
            parsed = json.loads(rawv) if rawv.startswith("[") else None
        except Exception:
            parsed = None
        parts = parsed if isinstance(parsed, list) else re.split(r"[,;|]\s*|\s{2,}", rawv)
        for m in parts:
            m = str(m).strip(" []\"'")
            if m:
                models[m] += 1
                providers[provider(m)] += 1

report["model_metadata"] = {
    "column": model_f, "rows_with_model_data": model_rows,
    "models": models.most_common(200), "providers": providers.most_common()
}

seen = defaultdict(list)
empty_content = 0
word_counts = []
variable_count = Counter()
var_re = re.compile(r"(\$\{[^}]+\}|\{\{[^}]+\}|\[\[[^]]+\]\])")

for i, r in enumerate(rows):
    c = (r.get(content_f) if content_f else "") or ""
    n = norm(c)
    if not n: empty_content += 1
    else: seen[n].append(i)
    variable_count[len(var_re.findall(c))] += 1
    word_counts.append(len(c.split()))

dupes = [v for v in seen.values() if len(v) > 1]
report["quality"] = {
    "empty_content_rows": empty_content,
    "exact_duplicate_groups": len(dupes),
    "rows_in_exact_duplicate_groups": sum(len(x) for x in dupes),
    "variable_count_distribution": variable_count.most_common(),
    "avg_words": round(sum(word_counts)/len(word_counts), 1) if word_counts else 0,
    "min_words": min(word_counts) if word_counts else 0,
    "max_words": max(word_counts) if word_counts else 0
}

for i, r in enumerate(rows[:20]):
    report["samples"].append({
        "row": i + 2,
        "title": (r.get(title_f) or "")[:200] if title_f else None,
        "category": (r.get(category_f) or "")[:100] if category_f else None,
        "type": (r.get(type_f) or "")[:50] if type_f else None,
        "content_preview": re.sub(r"\s+", " ", (r.get(content_f) or ""))[:300] if content_f else None
    })

Path("reports").mkdir(exist_ok=True)
Path("reports/prompt-analysis.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

md = [
    "# prompts.chat — Prompt Veri Analizi", "",
    f"- CSV satır sayısı: **{len(rows):,}**",
    f"- CSV boyutu: **{len(raw):,} bayt**"
]
if report["live_api"].get("status") == "ok":
    md += [
        f"- Canlı public API toplamı: **{report['live_api'].get('total'):,}**",
        f"- API toplam sayfa: **{report['live_api'].get('totalPages'):,}"
    ]
md += [f"- CSV sütunları: `{', '.join(fields)}`", "", "## Model metadata"]
if model_f:
    md += [f"- Model alanı: `{model_f}`", f"- Model bilgisi bulunan satır: **{model_rows:,}**", "- Sağlayıcı dağılımı:"]
    md += [f"  - {k}: {v:,}" for k,v in providers.most_common()]
else:
    md += ["- CSV snapshot içinde ayrı bir model/provider alanı bulunamadı."]
md += [
    "", "## Kalite / tekrar kontrolü",
    f"- Boş içerik: **{empty_content:,}**",
    f"- Birebir normalize edilmiş duplicate grubu: **{len(dupes):,}**",
    f"- Duplicate gruplarındaki satır: **{sum(len(x) for x in dupes):,}**",
    f"- Ortalama kelime: **{report['quality']['avg_words']}**",
    f"- Kelime aralığı: **{report['quality']['min_words']}–{report['quality']['max_words']}**",
    "", "## İlk 20 örnek"
]
for s in report["samples"]:
    md.append(f"- **{s['title']}** — {s['content_preview']}")
Path("reports/prompt-analysis.md").write_text("\n".join(md) + "\n", encoding="utf-8")

# Build a usable Excel workbook from the real snapshot.
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
except Exception as e:
    raise RuntimeError("openpyxl is required: " + str(e))

wb = Workbook()
ws = wb.active
ws.title = "PROMPTLAR"

headers = [
    "ID","Prompt Başlığı","Kategori","Etiketler","Prompt",
    "Kullanım Alanı","AI / Model","Sağlayıcı","Model Bağımsız",
    "Prompt Tipi","Değişkenler","Kelime Sayısı","Kalite Durumu",
    "Duplicate","Kaynak","Lisans"
]
ws.append(headers)

seen_index = {}
for i, r in enumerate(rows, start=1):
    content = (r.get(content_f) or "") if content_f else ""
    title = (r.get(title_f) or "") if title_f else ""
    category = (r.get(category_f) or "") if category_f else ""
    tags = (r.get(tags_f) or "") if tags_f else ""
    model_raw = (r.get(model_f) or "") if model_f else ""
    model_display = model_raw
    provs = set()
    if model_raw:
        try:
            parsed = json.loads(model_raw) if model_raw.startswith("[") else None
        except Exception:
            parsed = None
        parts = parsed if isinstance(parsed, list) else re.split(r"[,;|]\s*|\s{2,}", model_raw)
        clean = [str(x).strip(" []\"'") for x in parts if str(x).strip(" []\"'")]
        model_display = ", ".join(clean)
        provs = {provider(x) for x in clean}
    vars_found = var_re.findall(content)
    ncontent = norm(content)
    dup = "Evet" if ncontent and len(seen.get(ncontent, [])) > 1 else "Hayır"
    quality = "Boş" if not ncontent else ("Çok kısa" if len(content.split()) < 5 else "Normal")
    independent = "Belirtilmemiş" if not model_raw else "Hayır"
    if not model_raw:
        independent = "Belirtilmemiş"
    row_id = r.get(id_f) or i
    ws.append([
        clean_excel(row_id), clean_excel(title), clean_excel(category), clean_excel(tags), clean_excel(content),
        clean_excel(category), clean_excel(model_display), clean_excel(", ".join(sorted(provs))),
        clean_excel(independent), clean_excel((r.get(type_f) or "") if type_f else ""),
        clean_excel(", ".join(vars_found)), len(content.split()), clean_excel(quality), clean_excel(dup),
        "prompts.chat / prompts.csv", "CC0 1.0"
    ])

# Search sheet: search title, category, tags, prompt, use case, model, provider, type and variables.
search = wb.create_sheet("ARAMA")
search["A1"] = "PROMPT KATALOĞU — ARAMA"
search["A2"] = "Arama kelimesi:"
search["B2"] = ""
search["A3"] = "Örnek:"
search["B3"] = "video, ChatGPT, kod, analiz, Instagram"
search["A5"] = "Sonuçlar"
search["A6"] = "Arama; başlık, kategori, etiket, prompt, kullanım alanı, model, sağlayıcı, tip ve değişken alanlarını tarar."
search["A8"] = headers[0]
for j,h in enumerate(headers[1:], start=2):
    search.cell(8,j).value = h

last = len(rows) + 1
search_formula = f'=IF($B$2="","Arama kutusuna bir kelime yazın.",IFERROR(FILTER(PROMPTLAR!A2:P{last},ISNUMBER(SEARCH(LOWER($B$2),LOWER(PROMPTLAR!B2:B{last}&" "&PROMPTLAR!C2:C{last}&" "&PROMPTLAR!D2:D{last}&" "&PROMPTLAR!E2:E{last}&" "&PROMPTLAR!F2:F{last}&" "&PROMPTLAR!G2:G{last}&" "&PROMPTLAR!H2:H{last}&" "&PROMPTLAR!J2:J{last}&" "&PROMPTLAR!K2:K{last})))),"Sonuç bulunamadı."))'
search["A9"] = search_formula

# Info sheets
cats = wb.create_sheet("KATEGORİLER")
cats.append(["Kategori","Prompt Sayısı"])
for k,v in report["counts"].get("category", []): cats.append([k,v])

models_ws = wb.create_sheet("AI MODELLERİ")
models_ws.append(["Model","Kayıt Sayısı","Sağlayıcı"])
for k,v in models.most_common(): models_ws.append([k,v,provider(k)])

quality_ws = wb.create_sheet("KALİTE ANALİZİ")
quality_ws.append(["Metrik","Değer"])
for k,v in report["quality"].items():
    if isinstance(v,(int,float,str)): quality_ws.append([k,v])

info = wb.create_sheet("BİLGİ & LİSANS")
info.append(["Alan","Açıklama"])
info.append(["Veri kaynağı",CSV_URL])
info.append(["Snapshot satır sayısı",len(rows)])
info.append(["Canlı API toplamı",report["live_api"].get("total","Kontrol edilemedi")])
info.append(["Prompt içeriği lisansı","CC0 1.0 Universal — upstream LICENSE'a göre"])
info.append(["Kod/site lisansı","MIT — prompt verisiyle karıştırılmamalı"])
info.append(["Not","Model desteği yalnızca kaynakta açıkça belirtilmişse işaretlenmiştir; eksik bilgi tahmin edilmemiştir."])
info.append(["Not","Türkçe açıklama/çeviri alanı bu sürümde otomatik uydurulmamıştır; sonraki ürünleştirme aşamasında ayrı işlenecektir."])

# Styling
header_fill = PatternFill("solid", fgColor="1F4E78")
thin = Side(style="thin", color="D9E1F2")
for sheet in wb.worksheets:
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.sheet_view.showGridLines = False

ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:P{last}"
widths = [12,30,22,25,65,25,30,20,20,18,30,14,16,12,28,14]
for i,w in enumerate(widths,1):
    ws.column_dimensions[get_column_letter(i)].width = w
for row in ws.iter_rows(min_row=2, max_row=last):
    row[4].alignment = Alignment(wrap_text=True, vertical="top")
    for c in row:
        c.border = Border(bottom=thin)

search.column_dimensions["A"].width = 22
search.column_dimensions["B"].width = 80
search["B2"].font = Font(size=16, bold=True)
search["B2"].fill = PatternFill("solid", fgColor="FFF2CC")
search["A1"].font = Font(size=18, bold=True)
for c in search[8]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = header_fill
search.freeze_panes = "A9"

for sh in [cats, models_ws, quality_ws, info]:
    for col in range(1, sh.max_column+1):
        sh.column_dimensions[get_column_letter(col)].width = 28
    sh.freeze_panes = "A2"

out = Path("reports/prompts-chat-prompt-katalogu.xlsx")
wb.save(out)
print(f"Excel created: {out} ({out.stat().st_size:,} bytes)")
